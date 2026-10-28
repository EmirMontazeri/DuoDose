"""Native-screen ingestion, isolated fitting, provenance and complete-grid assessment."""

from __future__ import annotations

from dataclasses import asdict
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

import numpy as np
from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest
from rdkit import Chem

from reverse_dose.assessment import matrix_design
from reverse_dose.cli import artifact_paths, run
from reverse_dose.config import config_from_dict, load_config
from reverse_dose.data import DataBundle, DoseTransform, ExpressionTransform
from reverse_dose.downloads import fetch_file, select_pubchem_cid, verify_downloads
from reverse_dose.layout import DatasetPaths
from reverse_dose.preprocess import preprocess_sources
from reverse_dose.screen_metadata import map_screen_cells
from reverse_dose.screen_readers import read_almanac
from reverse_dose.utils import atomic_json, sha256_file


def _references(root: Path) -> None:
    directory = DatasetPaths(root, "Reference").raw / "cell_model_passports"
    directory.mkdir(parents=True)
    models = [f"m{i}" for i in range(20)]
    pd.DataFrame(
        {
            "model_id": models,
            "model_name": [f"C{i}" for i in range(20)],
            "COSMIC_ID": [str(100 + i) for i in range(20)],
        }
    ).to_csv(directory / "model_list_test.csv", index=False)
    pd.DataFrame(
        {"model_id": models, "RNASeq Broad Cell Lines": True, "RNASeq Sanger Cell Lines": False}
    ).to_csv(directory / "model_dataset_availability_test.csv", index=False)
    rows = ["model_id,gene_id,rsem_tpm"]
    for i, model in enumerate(models):
        for gene in range(12):
            rows.append(f"{model},g{gene},{1 + ((i + 1) * (gene + 3)) % 29}")
    with ZipFile(directory / "rnaseq_all_test.zip", "w") as archive:
        archive.writestr("expression.csv", "\n".join(rows) + "\n")


def _inventory(raw: Path, source: str) -> None:
    files = [
        {
            "path": p.relative_to(raw).as_posix(),
            "sha256": sha256_file(p),
            "url": "https://wiki.nci.nih.gov/fixture",
            "bytes": p.stat().st_size,
        }
        for p in sorted(raw.rglob("*"))
        if p.is_file() and p.suffix not in {".csv", ".xls"}
    ]
    atomic_json(
        raw / "official_downloads.json",
        {"schema_version": "official-downloads-v1", "source": source, "files": files},
    )


def _almanac(root: Path) -> tuple[Path, dict]:
    raw = DatasetPaths(root, "NCI-ALMANAC").raw
    raw.mkdir(parents=True)
    rows = []
    for cell in range(6):
        for i, a in enumerate([1e-8, 1e-7, 1e-6, 1e-5, 1e-4], start=1):
            for j, b in enumerate([2e-7, 2e-6, 2e-5], start=1):
                rows.append(
                    {
                        "COMBODRUGSEQ": len(rows) + 1,
                        "SCREENER": "1A",
                        "STUDY": "study",
                        "TESTDATE": "01/01/2017",
                        "PLATE": f"plate{j}",
                        "PANELNBR": cell // 3 + 1,
                        "CELLNBR": cell % 3 + 1,
                        "NSC1": 22,
                        "SAMPLE1": 1,
                        "NSC2": 11,
                        "SAMPLE2": 1,
                        "CONC1": a,
                        "CONC2": b,
                        "CONCUNIT1": "M",
                        "CONCUNIT2": "M",
                        "VALID": "Y",
                        "CELLNAME": f"C{cell}",
                        "PERCENTGROWTH": -50,
                        "PERCENTGROWTHNOTZ": np.nan if j == 1 else 40.0,
                        "TESTVALUE": 2.0,
                        "CONTROLVALUE": 4.0,
                    }
                )
    path = raw / "ComboDrugGrowth_Nov2017.csv"
    pd.DataFrame(rows).to_csv(path, index=False)
    with Chem.SDWriter(str(raw / "ComboCompoundSet.sdf")) as writer:
        for nsc, smiles in [(11, "CCO"), (22, "CCN")]:
            mol = Chem.MolFromSmiles(smiles)
            mol.SetProp("NSC", str(nsc))
            writer.write(mol)
    (raw / "ComboCompoundNames_small.txt").write_text("11\tEthanol\n22\tEthylamine\n")
    (raw / "ALMANAC_DataFields.txt").write_text("Official field definitions fixture")
    _inventory(raw, "nci_almanac")
    return path, {"sources": ["nci_almanac"], "workers": 1, "fingerprint_bits": 32}


def _oneil(root: Path) -> tuple[Path, dict]:
    raw = DatasetPaths(root, "ONeil-Merck").raw
    raw.mkdir(parents=True)
    path = raw / "response_w2lrww.xls"
    book = Workbook()
    sheet = book.active
    sheet.title = "combination response"
    sheet.append(
        [
            "BatchID",
            "cell_line",
            "drugA_name",
            "drugA Conc (µM)",
            "drugB_name",
            "drugB Conc (µM)",
            "combination_name",
            "viability1",
            "viability2",
            "viability3",
            "viability4",
        ]
    )
    for cell in range(6):
        for batch in (1, 3):
            for a in (0.01, 0.1, 1, 10):
                for b in (0.02, 0.2, 2, 20):
                    sheet.append(
                        [batch, f"C{cell}", "B", a, "A", b, "B & A", 0.5, "NULL", 4.0, 0.75]
                    )
    # O'Neil's file is an OOXML workbook despite the published .xls suffix.
    with path.open("wb") as handle:
        book.save(handle)
    book.close()
    for cid, label, smiles in [(1, "A", "CCO"), (2, "B", "CCN")]:
        directory = raw / "pubchem" / hashlib.sha256(label.encode()).hexdigest()[:16]
        atomic_json(directory / "identity.json", {"IdentifierList": {"CID": [cid]}})
        atomic_json(
            directory / "synonyms.json",
            {"InformationList": {"Information": [{"CID": cid, "Synonym": [label]}]}},
        )
        atomic_json(
            directory / "property.json",
            {"PropertyTable": {"Properties": [{"CID": cid, "SMILES": smiles}]}},
        )
    _inventory(raw, "oneil")
    return path, {"sources": ["oneil"], "workers": 1, "fingerprint_bits": 32}


@pytest.mark.parametrize(
    "source,dataset", [("oneil", "ONeil-Merck"), ("nci_almanac", "NCI-ALMANAC")]
)
def test_native_sources_preserve_coordinates_exclusions_and_verified_cache(
    tmp_path: Path, source: str, dataset: str
) -> None:
    root = tmp_path / "dataset"
    _references(root)
    path, settings = (_oneil if source == "oneil" else _almanac)(root)
    config = config_from_dict({"sources": settings})
    output = DatasetPaths(root, dataset).processed / "source"
    digest = sha256_file(path)
    manifest = preprocess_sources(
        root,
        asdict(config.sources),
        output,
        dataset_name=dataset,
        reference_dataset_name="Reference",
    )
    frame = pd.read_parquet(output / "observations.parquet")
    design = pd.read_parquet(output / "matrix_design.parquet")
    exclusions = pd.read_parquet(output / "exclusions.parquet")
    assert sha256_file(path) == digest
    assert frame.observation_id.is_unique
    assert frame.cell_id.nunique() == 6  # Includes repeated CELLNBR in different panels.
    assert (frame.drug_a.astype(str) < frame.drug_b.astype(str)).all()
    expected_a = np.array([0.02, 0.2, 2, 20] if source == "oneil" else [0.2, 2, 20])
    expected_b = np.array([0.01, 0.1, 1, 10] if source == "oneil" else [0.01, 0.1, 1, 10, 100])
    np.testing.assert_allclose(frame.dose_a_um, expected_a[frame.grid_col.to_numpy() - 1])
    np.testing.assert_allclose(frame.dose_b_um, expected_b[frame.grid_row.to_numpy() - 1])
    if source == "oneil":
        assert frame.matrix_id.nunique() == 24  # 6 cells × 2 batches × 2 usable replicates.
        assert frame.response_origin.unique().tolist() == ["viability1", "viability4"]
        assert set(exclusions.reason) == {"missing_response", "response_outside_quality_range"}
        assert (
            exclusions.loc[
                exclusions.reason.eq("response_outside_quality_range"), "response_percent"
            ]
            == 400
        ).all()
    else:
        assert frame.matrix_id.nunique() == 6
        assert frame.unit1.eq("M").all() and frame.unit2.eq("M").all()
        assert len(frame) == 90  # Complete 5×3 grids across three plates.
        assert manifest["counts"]["response_origins"]["reconstructed_test_over_control"] == 30
        assert set(frame.inhibition) == {0.5, 0.6}
    # The source cache validates artifacts and reuses them without parsing again.
    again = preprocess_sources(
        root,
        asdict(config.sources),
        output,
        dataset_name=dataset,
        reference_dataset_name="Reference",
    )
    assert again["identity_hash"] == manifest["identity_hash"]
    assert not (DatasetPaths(root, "Reference").processed).exists()
    assert set(design) == {
        "matrix_id",
        "cell_id",
        "drug_a",
        "drug_b",
        "dose_a_um",
        "dose_b_um",
        "grid_row",
        "grid_col",
        "observation_id",
    }
    (output / "matrix_design.parquet").write_bytes(b"corrupted")
    with pytest.raises(ValueError, match="integrity"):
        preprocess_sources(
            root,
            asdict(config.sources),
            output,
            dataset_name=dataset,
            reference_dataset_name="Reference",
        )


def test_almanac_endpoint_policy_never_substitutes_time_zero_growth(tmp_path: Path) -> None:
    path, settings = _almanac(tmp_path)
    config = config_from_dict({"sources": settings}).sources
    frame, _ = read_almanac(path, "raw/source.csv", asdict(config))
    assert set(frame.response_percent) == {40, 50}
    config.nci_missing_no_tz = "exclude"
    frame, _ = read_almanac(path, "raw/source.csv", asdict(config))
    assert frame.response_percent.isna().sum() == 30
    config.nci_missing_no_tz = "error"
    with pytest.raises(ValueError, match="PERCENTGROWTHNOTZ"):
        read_almanac(path, "raw/source.csv", asdict(config))
    config.nci_response = "percent_growth"
    frame, _ = read_almanac(path, "raw/source.csv", asdict(config))
    assert frame.response_percent.eq(-50).all()
    bad = pd.read_csv(path)
    bad.loc[0, "CONCUNIT1"] = "mg/mL"
    bad.to_csv(path, index=False)
    with pytest.raises(ValueError, match="Unsupported concentration unit"):
        read_almanac(path, "raw/source.csv", asdict(config))


def test_cell_derivatives_missing_expression_and_alias_conflicts(tmp_path: Path) -> None:
    _references(tmp_path)
    reference = DatasetPaths(tmp_path, "Reference").raw
    p = reference / "cell_model_passports/model_list_test.csv"
    m = pd.read_csv(p)
    m.loc[1, "model_name"] = "C0_derivative"
    m["CCLE_ID"] = ["C0_TISSUE", "C0_TISSUE", *[f"C{i}_TISSUE" for i in range(2, 20)]]
    m.to_csv(p, index=False)
    a = reference / "cell_model_passports/model_dataset_availability_test.csv"
    av = pd.read_csv(a)
    av.loc[2, "RNASeq Broad Cell Lines"] = False
    av.to_csv(a, index=False)
    cells, names, provenance, _ = map_screen_cells(
        ["C0", "C0_derivative", "C2", "UNKNOWN"], reference, {}, "exclude"
    )
    assert names["C0"] == "m0" and names["C0_derivative"] == "m1"
    assert set(cells.model_id) == {"m0", "m1"}
    assert provenance["excluded_labels"]["C2"]["reason"] == "missing_cell_expression"
    with pytest.raises(ValueError, match="conflicts"):
        map_screen_cells(
            ["C0"], reference, {"C0": {"cell_id": "m1", "evidence": "not valid"}}, "exclude"
        )
    with pytest.raises(ValueError, match="UNKNOWN"):
        map_screen_cells(["UNKNOWN"], reference, {}, "error")


@pytest.mark.parametrize(
    "source,dataset", [("oneil", "ONeil-Merck"), ("nci_almanac", "NCI-ALMANAC")]
)
def test_independent_training_evaluation_export_contract(
    tmp_path: Path, source: str, dataset: str
) -> None:
    root = tmp_path / "dataset"
    _references(root)
    _, sources = (_oneil if source == "oneil" else _almanac)(root)
    config = config_from_dict(
        {
            "dataset_root": str(root),
            "dataset_name": dataset,
            "cell_reference_dataset": "Reference",
            "sources": sources,
            "features": {"n_genes": 8},
            "threads": 1,
            "vae": {"hidden_dims": [8], "latent_dim": 3},
            "model": {
                "encoder_dim": 4,
                "encoder_hidden_dim": 8,
                "mf_dim": 2,
                "deep_hidden_dims": [8],
                "deep_dim": 8,
                "mixture_components": 2,
            },
            "vae_training": {"epochs": 1, "batch_size": 16},
            "training": {"epochs": 1, "batch_size": 64},
            "evaluation": {"bootstrap_repetitions": 100},
        }
    )
    result = run(config, tmp_path / "run")
    assert result["status"] == "complete"
    bundle = DataBundle(artifact_paths(config)[2])
    assert "matrix_design.parquet" in bundle.manifest["artifacts"]
    assert (
        bundle.manifest["features"]["inhibition_definition"]
        == json.loads((bundle.path / "source_manifest.json").read_text())["response_definition"]
    )
    design = matrix_design(bundle, root, "test", dataset_name=dataset)
    assert set(design.matrix_id) == set(
        bundle.observations.loc[bundle.observations.split.eq("test"), "matrix_id"]
    )
    metrics = json.loads((tmp_path / "run/evaluation/test/metrics.json").read_text())
    assert metrics["target_achievement"]["matrices"] > 0
    assert not (DatasetPaths(root, "Reference").splits).exists()
    from reverse_dose.inference import Predictor, export_model

    exported = tmp_path / "export"
    export_model(bundle.path, Path(result["model_checkpoint"]), exported, root)
    predictor = Predictor.from_export(exported)
    assert predictor.drug_lookup
    # Assessment continues to work from a verified portable design without raw files.
    assert len(matrix_design(bundle, tmp_path / "absent_raw", "test", dataset_name=dataset)) == len(
        design
    )


def test_full_design_keeps_missing_and_qc_excluded_coordinates(tmp_path: Path) -> None:
    root = tmp_path / "dataset"
    _references(root)
    path, settings = _oneil(root)
    with path.open("rb") as handle:
        workbook = load_workbook(handle)
    workbook.active["H2"] = "NULL"
    workbook.active["H3"] = 4.0
    with path.open("wb") as handle:
        workbook.save(handle)
    workbook.close()
    config = config_from_dict({"sources": settings})
    output = DatasetPaths(root, "ONeil-Merck").processed / "source"
    preprocess_sources(
        root,
        asdict(config.sources),
        output,
        dataset_name="ONeil-Merck",
        reference_dataset_name="Reference",
    )
    rows = pd.read_parquet(output / "observations.parquet")
    design = pd.read_parquet(output / "matrix_design.parquet")
    exclusions = pd.read_parquet(output / "exclusions.parquet")
    points = exclusions.loc[exclusions.replicate.eq(1)]
    assert len(points) == 2
    for point in points.itertuples():
        expected = f"{point.matrix_id}:{point.grid_row}:{point.grid_col}"
        assert expected in set(design.observation_id)
        assert expected not in set(rows.observation_id)


def test_external_evaluation_reuses_training_state_and_reports_exposure(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from reverse_dose.transfer import ExternalBundle, evaluate_external

    root = tmp_path / "dataset"
    _references(root)
    _, sources = _almanac(root)
    values = {
        "dataset_root": str(root),
        "dataset_name": "NCI-ALMANAC",
        "cell_reference_dataset": "Reference",
        "sources": sources,
        "features": {"n_genes": 8},
        "threads": 1,
        "vae": {"hidden_dims": [8], "latent_dim": 3},
        "model": {
            "encoder_dim": 4,
            "encoder_hidden_dim": 8,
            "mf_dim": 2,
            "deep_hidden_dims": [8],
            "deep_dim": 8,
            "mixture_components": 2,
        },
        "vae_training": {"epochs": 1, "batch_size": 16},
        "training": {"epochs": 1, "batch_size": 64},
        "evaluation": {"bootstrap_repetitions": 100},
    }
    config = config_from_dict(values)
    result = run(config, tmp_path / "trained")
    training = DataBundle(Path(result["data_dir"]))
    from reverse_dose.training import train_stage

    wrong = config_from_dict(config.to_dict())
    wrong.dataset_name = "Wrong Dataset"
    with pytest.raises(ValueError, match="dataset names differ"):
        train_stage(training, wrong, tmp_path / "wrong_run", "vae")
    assert not (tmp_path / "wrong_run").exists()
    checkpoint = Path(result["model_checkpoint"])
    before = sha256_file(checkpoint)
    _, target_sources = _oneil(root)
    target_config = config_from_dict(
        {**values, "dataset_name": "ONeil-Merck", "sources": target_sources}
    )

    def forbidden_fit(*args, **kwargs):
        raise AssertionError("External evaluation must never fit expression or dose state")

    monkeypatch.setattr(ExpressionTransform, "fit", forbidden_fit)
    monkeypatch.setattr(DoseTransform, "fit", forbidden_fit)
    report = evaluate_external(training.path, checkpoint, target_config, tmp_path / "external")
    assert report["generalization"].startswith("external-screen")
    assert report["external_protocol"]["target_cells_seen_in_supervised_training"] == 4
    assert report["external_protocol"]["target_cells"] == 6
    assert sha256_file(checkpoint) == before
    assert report["dataset_fingerprint"] != training.manifest["fingerprint"]
    source, expression, _ = artifact_paths(target_config)
    target = ExternalBundle(training, source, expression, checkpoint, "unseen_supervised")
    assert len(target.cell_ids) == 2
    assert not target.observations.seen_in_supervised_training.any()
    assert target.dose_transform is training.dose_transform
    assert target.expression_transform is training.expression_transform
    assert set(target.drug_ids) == {"A", "B"}
    assert not artifact_paths(target_config)[2].exists()
    manifest_path = source / "source_manifest.json"
    native = json.loads(manifest_path.read_text())
    native["fingerprint_definition"]["radius"] += 1
    atomic_json(manifest_path, native)
    with pytest.raises(ValueError, match="fingerprint definitions differ"):
        ExternalBundle(training, source, expression, checkpoint)


def test_official_downloads_refuse_overwrites_and_tampering(tmp_path: Path) -> None:
    item = fetch_file(tmp_path, "record.txt", "https://wiki.nci.nih.gov/test", payload=b"original")
    assert item["sha256"] == hashlib.sha256(b"original").hexdigest()
    with pytest.raises(ValueError, match="immutable"):
        fetch_file(tmp_path, "record.txt", "https://wiki.nci.nih.gov/test", payload=b"changed")
    atomic_json(
        tmp_path / "official_downloads.json",
        {"schema_version": "official-downloads-v1", "files": [item]},
    )
    verify_downloads(tmp_path)
    (tmp_path / "record.txt").write_bytes(b"changed")
    with pytest.raises(ValueError, match="integrity"):
        verify_downloads(tmp_path)


def test_ambiguous_chemistry_only_accepts_identical_inorganic_salt_parent(tmp_path: Path) -> None:
    def record(cid: int, smiles: str) -> None:
        atomic_json(
            tmp_path / f"candidate_{cid}.json",
            {"PropertyTable": {"Properties": [{"CID": cid, "SMILES": smiles}]}},
        )

    record(1, "CCN.Cl.Cl")
    record(2, "CCN")
    assert select_pubchem_cid(tmp_path, [1, 2]) == 2
    record(1, "CCO.Cl.Cl")
    with pytest.raises(ValueError, match="not unique"):
        select_pubchem_cid(tmp_path, [1, 2])


@pytest.mark.parametrize(
    "values",
    [
        {"sources": {"sources": ["dream", "oneil"]}},
        {"sources": {"sources": ["nci_almanac", "oneil"]}},
        {"sources": {"nci_response": "synergy_score"}},
        {"sources": {"nci_missing_no_tz": "use_growth"}},
        {"sources": {"oneil_batches": [1, 1]}},
        {"sources": {"oneil_batches": [4]}},
        {"cell_reference_dataset": "../outside"},
    ],
)
def test_invalid_screen_configuration_fails_early(values: dict) -> None:
    with pytest.raises(ValueError):
        config_from_dict(values)


@pytest.mark.parametrize(
    "path",
    [
        "configs/nci_almanac.yaml",
        "configs/oneil.yaml",
        "configs/nci_almanac_smoke.yaml",
        "configs/oneil_smoke.yaml",
    ],
)
def test_shipped_screen_configurations_remain_independent(path: str) -> None:
    config = load_config(Path(path))
    assert len(config.sources.sources) == 1
    assert config.cell_reference_dataset == "AstraZeneca–DREAM"
    assert config_from_dict(config.to_dict()) == config
    assert all(config.dataset_name in p.parts for p in artifact_paths(config))
