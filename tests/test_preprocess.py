"""Regression tests for immutable source ingestion and scientific orientation."""

from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path
import re
from zipfile import ZipFile

import numpy as np
from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest

from reverse_dose.preprocess import (
    INHIBITION_QUALITY_MAX,
    INHIBITION_QUALITY_MIN,
    INHIBITION_QUALITY_REASON,
    build_drug_features,
    concentration_factor_um,
    map_cells,
    parse_corrected_matrix,
    preprocess_sources,
    resolve_pubchem_cid,
    resolve_smiles,
    response_quality_control,
)
from reverse_dose.layout import DatasetPaths


def write_matrix(
    path: Path,
    *,
    title: str = "Cell1",
    response: float | str | None = 60.0,
    bad_dimensions: bool = False,
) -> None:
    """Agent B lies on rows in nM; Agent A lies on columns in micromolar."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    for row in [
        ["NA", 0, 2, "(=Agent 2)"],
        [0, 100, 80, None],
        [1000, 70, response, None],
        ["(=Agent 1)", None, None, None],
        ["Agent1", "B"],
        ["Agent2", "A"],
        ["Unit1", "nM"],
        ["Unit2", "\\\\muM"],
        ["Title", title],
    ]:
        worksheet.append(row)
    workbook.save(path)
    if bad_dimensions:
        with ZipFile(path) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        members["xl/worksheets/sheet1.xml"] = re.sub(
            rb'<dimension ref="[^"]+"',
            b'<dimension ref="A1"',
            members["xl/worksheets/sheet1.xml"],
        )
        with ZipFile(path, "w") as archive:
            for name, content in members.items():
                archive.writestr(name, content)


@pytest.fixture
def raw_dataset(tmp_path: Path) -> Path:
    root = tmp_path / "dataset"
    az = root / "raw/AstraZeneca–DREAM/az_dream"
    molecular = root / "raw/AstraZeneca–DREAM/az_dream_molecular"
    cmp = root / "raw/AstraZeneca–DREAM/cell_model_passports"
    for directory in [az, molecular, cmp]:
        directory.mkdir(parents=True)
    pd.DataFrame(
        {
            "Sanger.Name": ["Cell1"],
            "Alternative.Name": ["Cell-1"],
            "CCLE.Name": ["CELL1_LUNG"],
            "COSMIC": [123],
        }
    ).to_csv(molecular / "cell_info.csv", index=False)
    pd.DataFrame({"model_id": ["SIDM0001"], "COSMIC_ID": ["123.0"]}).to_csv(
        cmp / "model_list_20260101.csv", index=False
    )
    drugs = pd.DataFrame(
        {
            "ChallengeName": ["A", "B", "Unsupported"],
            "SMILES or PubChem ID": ["CCO", "CCN", None],
        }
    )
    drugs.to_csv(az / "Drug_info_release.csv", index=False)
    drugs.to_excel(
        az / "Extra DrugChemistryData to release for Nat Comms FINAL RELEASE_200219.xlsx",
        sheet_name="Drug_Chemistry_Data",
        index=False,
    )
    (az / "Combenefit_dose_response_fitting_code.txt").write_text("EC0=100.0\n")
    write_matrix(
        az / "DREAM_Challenge_matrices_corrected/A&B&Cell1&Exp1.xlsx",
        bad_dimensions=True,
    )
    return root


def test_read_complete_grid_with_incorrect_declared_dimensions(tmp_path: Path) -> None:
    path = tmp_path / "A&B&Cell1&Exp1.xlsx"
    write_matrix(path, bad_dimensions=True)
    matrix = parse_corrected_matrix(path)
    assert matrix["declared_dimensions"] in {"A1", "A1:A1"}
    assert matrix["dose1_um"] == [0.0, 1.0]
    assert matrix["dose2_um"] == [0.0, 2.0]
    assert matrix["points"][-1] == (2, 2, 60.0)
    assert matrix["metadata"]["Agent1"] == "B"


@pytest.mark.parametrize(
    ("unit", "expected"),
    [
        ("M", 1e6),
        ("mM", 1e3),
        ("\\\\muM", 1),
        ("µM", 1),
        ("μM", 1),
        ("nM", 1e-3),
        ("pM", 1e-6),
    ],
)
def test_explicit_molar_unit_conversion(unit: str, expected: float) -> None:
    assert concentration_factor_um(unit) == expected


def test_mass_units_never_guessed() -> None:
    with pytest.raises(ValueError, match="Unsupported concentration unit"):
        concentration_factor_um("mg/mL")


@pytest.mark.parametrize("response", ["nonsense", float("inf"), True])
def test_malformed_response_fails(tmp_path: Path, response: float | str) -> None:
    path = tmp_path / "A&B&Cell1&Exp1.xlsx"
    # XLSX writers cannot store infinity: direct string infinity tests the parser.
    write_matrix(path, response="inf" if response == float("inf") else response)
    with pytest.raises(ValueError):
        parse_corrected_matrix(path)


def test_missing_response_is_explicit(tmp_path: Path) -> None:
    path = tmp_path / "A&B&Cell1&Exp1.xlsx"
    write_matrix(path, response=None)
    matrix = parse_corrected_matrix(path)
    assert matrix["missing_points"] == [(2, 2)]
    assert len(matrix["points"]) == 3


def test_structure_alternatives_require_verified_equivalence() -> None:
    canonical, status = resolve_smiles("CCO;OCC")
    assert canonical == "CCO"
    assert status == "equivalent_semicolon_smiles"
    assert resolve_smiles("CCO;CCN") == (None, "ambiguous_semicolon_structures")
    assert resolve_smiles("12345") == (None, "pubchem_only_no_local_structure")
    assert resolve_smiles("notasmiles")[0] is None


def test_atomic_preprocessing_orientation_cache_and_raw_integrity(
    raw_dataset: Path,
) -> None:
    raw_hashes = {
        str(path): hashlib.sha256(path.read_bytes()).hexdigest()
        for path in (raw_dataset / "raw").rglob("*")
        if path.is_file()
    }
    destination = raw_dataset / "processed/AstraZeneca–DREAM/source"
    config = {"workers": 1, "fingerprint_bits": 32}
    manifest = preprocess_sources(raw_dataset, config, destination)
    frame = pd.read_parquet(destination / "observations.parquet")
    combination = frame.query("grid_row == 2 and grid_col == 2").iloc[0]
    assert combination.drug_a == "A" and combination.drug_b == "B"
    assert combination.dose_a_um == 2 and combination.dose_b_um == 1
    assert combination.inhibition == pytest.approx(0.4)
    assert combination.canonical_order_swapped
    assert manifest["counts"]["incorrect_declared_dimensions_read_completely"] == 1
    with np.load(destination / "drugs.npz", allow_pickle=False) as arrays:
        assert arrays["features"].shape == (2, 32)
        assert arrays["features"].dtype == np.float32
    assert (
        preprocess_sources(raw_dataset, config, destination)["identity_hash"]
        == manifest["identity_hash"]
    )
    assert raw_hashes == {
        name: hashlib.sha256(Path(name).read_bytes()).hexdigest() for name in raw_hashes
    }
    with pytest.raises(ValueError, match="differ"):
        preprocess_sources(raw_dataset, {**config, "include_zero_doses": False}, destination)
    (destination / "drugs.npz").write_bytes(b"corrupted")
    with pytest.raises(ValueError, match="integrity"):
        preprocess_sources(raw_dataset, config, destination)


def test_preprocessing_uses_only_the_selected_dataset(raw_dataset: Path) -> None:
    original = DatasetPaths(raw_dataset)
    selected = DatasetPaths(raw_dataset, "Another Dataset")
    original.raw.rename(selected.raw)
    output = selected.processed / "source"
    manifest = preprocess_sources(
        raw_dataset,
        {"workers": 1, "resolve_pubchem": False},
        output,
        dataset_name=selected.name,
    )
    assert manifest["raw_files"]
    for entry in manifest["raw_files"]:
        assert entry["path"].startswith("raw/Another Dataset/")
        assert (raw_dataset / entry["path"]).is_file()
    with pytest.raises(ValueError, match="No corrected dream"):
        preprocess_sources(raw_dataset, {"workers": 1}, original.processed / "source")


def test_source_cell_alias_requires_evidence(raw_dataset: Path) -> None:
    cells, names, provenance = map_cells(raw_dataset)
    assert cells.iloc[0].model_id == "SIDM0001"
    assert names["CELL1"] == "Cell1"
    assert "KMSII" not in names
    with pytest.raises(ValueError, match="evidence"):
        map_cells(raw_dataset, {"KMSII": {"cell_id": "Cell1", "evidence": ""}})
    _, names, provenance = map_cells(
        raw_dataset,
        {"Alias": {"cell_id": "Cell1", "evidence": "Synthetic test mapping"}},
    )
    assert names["ALIAS"] == "Cell1"
    assert provenance["aliases"]["Alias"]["evidence"] == "Synthetic test mapping"


def test_unresolved_and_conflicting_cells_quarantined(raw_dataset: Path) -> None:
    write_matrix(
        raw_dataset
        / "raw/AstraZeneca–DREAM/az_dream/DREAM_Challenge_matrices_corrected/A&B&Cell1&Exp2.xlsx",
        title="Cell1C",
    )
    destination = raw_dataset / "processed/AstraZeneca–DREAM/quarantine"
    manifest = preprocess_sources(raw_dataset, {"workers": 1}, destination)
    assert manifest["counts"]["retained_matrices"] == 1
    exclusions = pd.read_parquet(destination / "exclusions.parquet")
    assert set(exclusions.reason) == {"unresolved_cell_label", "cell_label_mismatch"}
    assert exclusions.source_file.str.endswith("Exp2.xlsx").all()
    with pytest.raises(ValueError, match="unresolved or conflicting"):
        preprocess_sources(
            raw_dataset,
            {"workers": 1, "unresolved_cell_policy": "error"},
            raw_dataset / "processed/AstraZeneca–DREAM/error",
        )


def test_missing_and_zero_doses_generate_explicit_exclusions(raw_dataset: Path) -> None:
    write_matrix(
        raw_dataset
        / "raw/AstraZeneca–DREAM/az_dream/DREAM_Challenge_matrices_corrected/A&B&Cell1&Exp2.xlsx",
        response=None,
    )
    destination = raw_dataset / "processed/AstraZeneca–DREAM/no_zeros"
    manifest = preprocess_sources(
        raw_dataset, {"workers": 1, "include_zero_doses": False}, destination
    )
    assert manifest["counts"]["retained_observations"] == 1
    assert manifest["counts"]["excluded_by_reason"] == {
        "missing_response": 1,
        "zero_dose_excluded_by_config": 6,
    }
    with pytest.raises(ValueError, match="missing response points"):
        preprocess_sources(
            raw_dataset,
            {"workers": 1, "missing_response_policy": "error"},
            raw_dataset / "processed/AstraZeneca–DREAM/missing_error",
        )


def test_raw_output_protection(raw_dataset: Path) -> None:
    with pytest.raises(ValueError, match="immutable"):
        preprocess_sources(raw_dataset, {}, raw_dataset / "raw/AstraZeneca–DREAM/generated")


def test_duplicate_cosmic_mappings_fail(raw_dataset: Path) -> None:
    path = raw_dataset / "raw/AstraZeneca–DREAM/cell_model_passports/model_list_20260101.csv"
    pd.DataFrame({"model_id": ["SIDM1", "SIDM2"], "COSMIC_ID": [123, 123]}).to_csv(
        path, index=False
    )
    with pytest.raises(ValueError, match="Ambiguous CMP"):
        map_cells(raw_dataset)


def test_response_outside_unit_interval_but_inside_quality_range_is_not_clipped(
    raw_dataset: Path,
) -> None:
    write_matrix(
        raw_dataset
        / "raw/AstraZeneca–DREAM/az_dream/DREAM_Challenge_matrices_corrected/A&B&Cell1&Exp1.xlsx",
        response=130,
    )
    destination = raw_dataset / "processed/AstraZeneca–DREAM/out_of_range"
    manifest = preprocess_sources(raw_dataset, {"workers": 1}, destination)
    assert manifest["counts"]["retained_observations_outside_unit_interval"] == 1
    assert pd.read_parquet(destination / "observations.parquet").inhibition.min() == pytest.approx(
        -0.3
    )
    assert json.loads((destination / "source_manifest.json").read_text())["config"][
        "include_zero_doses"
    ]


@pytest.mark.parametrize(
    ("response", "expected_inhibition"),
    [(300.0, INHIBITION_QUALITY_MIN), (-100.0, INHIBITION_QUALITY_MAX)],
)
def test_inhibition_quality_boundaries_are_inclusive(
    raw_dataset: Path, response: float, expected_inhibition: float
) -> None:
    path = (
        raw_dataset
        / "raw/AstraZeneca–DREAM/az_dream/DREAM_Challenge_matrices_corrected/A&B&Cell1&Exp1.xlsx"
    )
    write_matrix(path, response=response)
    destination = raw_dataset / "processed/AstraZeneca–DREAM/quality_boundary"
    manifest = preprocess_sources(raw_dataset, {"workers": 1}, destination)
    observations = pd.read_parquet(destination / "observations.parquet")

    assert any(value == pytest.approx(expected_inhibition) for value in observations.inhibition)
    assert manifest["response_quality_control"] == response_quality_control()
    assert manifest["counts"]["quality_control_excluded_observations"] == 0


@pytest.mark.parametrize(
    ("response", "expected_inhibition"),
    [(301.0, -2.01), (-101.0, 2.01)],
)
def test_inhibition_outside_quality_range_is_removed_and_preserved(
    raw_dataset: Path, response: float, expected_inhibition: float
) -> None:
    path = (
        raw_dataset
        / "raw/AstraZeneca–DREAM/az_dream/DREAM_Challenge_matrices_corrected/A&B&Cell1&Exp1.xlsx"
    )
    write_matrix(path, response=response)
    destination = raw_dataset / "processed/AstraZeneca–DREAM/quality_exclusion"
    manifest = preprocess_sources(raw_dataset, {"workers": 1}, destination)
    observations = pd.read_parquet(destination / "observations.parquet")
    exclusions = pd.read_parquet(destination / "exclusions.parquet")
    quality = exclusions.loc[exclusions.reason.eq(INHIBITION_QUALITY_REASON)]

    assert len(observations) == 3
    assert observations.inhibition.between(
        INHIBITION_QUALITY_MIN, INHIBITION_QUALITY_MAX, inclusive="both"
    ).all()
    assert len(quality) == 1
    assert quality.iloc[0].response_percent == pytest.approx(response)
    assert quality.iloc[0].inhibition == pytest.approx(expected_inhibition)
    assert "without clipping" in quality.iloc[0].detail
    assert manifest["exclusion_schema"][-2:] == ["response_percent", "inhibition"]
    assert manifest["counts"]["quality_control_excluded_observations"] == 1
    assert manifest["counts"]["excluded_by_reason"][INHIBITION_QUALITY_REASON] == 1


def test_pubchem_exact_cid_cache_is_verified_and_offline(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    payload = json.dumps(
        {"PropertyTable": {"Properties": [{"CID": 702, "SMILES": "CCO"}]}}
    ).encode()
    calls = []

    def fetch(request: object, timeout: int) -> io.BytesIO:
        calls.append((request, timeout))
        return io.BytesIO(payload)

    monkeypatch.setattr("reverse_dose.preprocess.urllib.request.urlopen", fetch)
    monkeypatch.setattr("reverse_dose.preprocess.time.sleep", lambda _: None)
    smiles, manifest, files = resolve_pubchem_cid(702, tmp_path)
    assert smiles == "CCO"
    assert manifest["requested_cid"] == 702
    assert manifest["response_sha256"] == hashlib.sha256(payload).hexdigest()
    assert len(calls) == 1
    assert resolve_pubchem_cid(702, tmp_path)[0] == "CCO"
    assert len(calls) == 1  # Existing cache never performs network I/O.
    files[0].write_bytes(b"corrupt")
    with pytest.raises(ValueError, match="integrity"):
        resolve_pubchem_cid(702, tmp_path)
    assert len(calls) == 1


def test_pubchem_rejects_wrong_cid_and_never_commits(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    payload = json.dumps(
        {"PropertyTable": {"Properties": [{"CID": 999, "SMILES": "CCO"}]}}
    ).encode()
    monkeypatch.setattr(
        "reverse_dose.preprocess.urllib.request.urlopen",
        lambda *args, **kwargs: io.BytesIO(payload),
    )
    with pytest.raises(ValueError, match="exact requested CID"):
        resolve_pubchem_cid(702, tmp_path)
    assert not (tmp_path / "702").exists()


def test_pubchem_release_precedence_and_explicit_offline_mode(
    raw_dataset: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    primary_path = raw_dataset / "raw/AstraZeneca–DREAM/az_dream/Drug_info_release.csv"
    drugs = pd.read_csv(primary_path)
    drugs.loc[drugs.ChallengeName == "Unsupported", "SMILES or PubChem ID"] = "702"
    drugs.to_csv(primary_path, index=False)
    local, _, _, _ = build_drug_features(raw_dataset, 32, 2, resolve_pubchem=False)
    assert not local.set_index("drug_id").loc["Unsupported", "supported"]
    payload = json.dumps(
        {"PropertyTable": {"Properties": [{"CID": 702, "SMILES": "CCO"}]}}
    ).encode()
    monkeypatch.setattr(
        "reverse_dose.preprocess.urllib.request.urlopen",
        lambda *args, **kwargs: io.BytesIO(payload),
    )
    monkeypatch.setattr("reverse_dose.preprocess.time.sleep", lambda _: None)
    resolved, _, _, paths = build_drug_features(raw_dataset, 32, 2, resolve_pubchem=True)
    row = resolved.set_index("drug_id").loc["Unsupported"]
    assert row.supported
    assert row.pubchem_cid == 702
    assert row.structure_status == "resolved_exact_released_pubchem_cid"
    assert any(path.endswith("702/response.json") for path in paths)
    assert row.pubchem_cache_manifest == (
        "processed/AstraZeneca–DREAM/drug_features/pubchem/v1/702/manifest.json"
    )
    assert (raw_dataset / row.pubchem_cache_manifest).is_file()


def test_released_gemcitibine_alias_has_exact_recorded_evidence(
    raw_dataset: Path,
) -> None:
    az = raw_dataset / "raw/AstraZeneca–DREAM/az_dream"
    primary_path = az / "Drug_info_release.csv"
    frame = pd.read_csv(primary_path)
    frame.loc[frame.ChallengeName == "B", "ChallengeName"] = "Gemcitabine"
    frame.to_csv(primary_path, index=False)
    frame["Drug name"] = [None, "Gemcitabine, Gemcitibine, Gemzar", None]
    frame.to_excel(
        az / "Extra DrugChemistryData to release for Nat Comms FINAL RELEASE_200219.xlsx",
        sheet_name="Drug_Chemistry_Data",
        index=False,
    )
    old = az / "DREAM_Challenge_matrices_corrected/A&B&Cell1&Exp1.xlsx"
    workbook = load_workbook(old)
    workbook.active["B5"] = "Gemcitibine"
    path = old.with_name("A&Gemcitabine&Cell1&Exp1.xlsx")
    workbook.save(path)
    old.unlink()
    destination = raw_dataset / "processed/AstraZeneca–DREAM/alias"
    manifest = preprocess_sources(raw_dataset, {"workers": 1}, destination)
    assert manifest["drug_alias_evidence"]["Gemcitibine"]["drug_id"] == "Gemcitabine"
    frame = pd.read_parquet(destination / "observations.parquet")
    assert set(frame.drug_b) == {"Gemcitabine"}
    assert set(frame.agent1) == {"Gemcitibine"}


def test_source_regeneration_is_byte_reproducible(raw_dataset: Path) -> None:
    first = raw_dataset / "processed/AstraZeneca–DREAM/first"
    second = raw_dataset / "processed/AstraZeneca–DREAM/second"
    a = preprocess_sources(raw_dataset, {"workers": 1}, first)
    b = preprocess_sources(raw_dataset, {"workers": 1}, second)
    assert a == b
    assert (first / "source_manifest.json").read_bytes() == (
        second / "source_manifest.json"
    ).read_bytes()


def test_entire_missing_oi_grid_is_explicitly_excluded(raw_dataset: Path) -> None:
    oi_dir = raw_dataset / "raw/AstraZeneca–DREAM/az_dream/DREAM_OI_matrices_corrected"
    valid_path = oi_dir / "A&B&Cell1&Exp1.xlsx"
    missing_path = oi_dir / "A&B&Cell1&Exp2.xlsx"
    write_matrix(valid_path)
    write_matrix(missing_path)
    workbook = load_workbook(missing_path)
    for row in workbook.active.iter_rows(min_row=2, max_row=3, min_col=2, max_col=3):
        for cell in row:
            cell.value = None
    workbook.save(missing_path)
    parsed = parse_corrected_matrix(missing_path)
    assert parsed["points"] == []
    assert parsed["missing_points"] == [(1, 1), (1, 2), (2, 1), (2, 2)]
    destination = raw_dataset / "processed/AstraZeneca–DREAM/oi"
    manifest = preprocess_sources(raw_dataset, {"workers": 1, "sources": ["oi"]}, destination)
    assert manifest["counts"]["parsed_matrices"] == 2
    assert manifest["counts"]["retained_matrices"] == 1
    assert manifest["counts"]["excluded_by_reason"] == {
        "missing_response": 4,
        "no_measured_responses": 1,
    }
    exclusions = pd.read_parquet(destination / "exclusions.parquet")
    assert set(exclusions.source) == {"oi"}
    assert exclusions.source_file.str.endswith("Exp2.xlsx").all()
    with pytest.raises(ValueError, match="missing response points"):
        preprocess_sources(
            raw_dataset,
            {"workers": 1, "sources": ["oi"], "missing_response_policy": "error"},
            raw_dataset / "processed/AstraZeneca–DREAM/oi_error",
        )


def test_source_manifest_integrity_and_complete_output_inventory(
    raw_dataset: Path,
) -> None:
    destination = raw_dataset / "processed/AstraZeneca–DREAM/integrity"
    preprocess_sources(raw_dataset, {"workers": 1}, destination)
    path = destination / "source_manifest.json"
    original = path.read_text()
    manifest = json.loads(original)
    manifest["counts"]["retained_observations"] = 999
    path.write_text(json.dumps(manifest))
    with pytest.raises(ValueError, match="manifest integrity"):
        preprocess_sources(raw_dataset, {"workers": 1}, destination)
    manifest = json.loads(original)
    del manifest["output_hashes"]["observations.parquet"]
    path.write_text(json.dumps(manifest))
    (destination / "source_manifest.sha256").write_text(
        hashlib.sha256(path.read_bytes()).hexdigest() + "\n"
    )
    with pytest.raises(ValueError, match="hash inventory"):
        preprocess_sources(raw_dataset, {"workers": 1}, destination)
