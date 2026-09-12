"""Immutable source ingestion for the explicit V2 experimental baseline.

Corrected matrices supply measured grid points, without interpolation or replicate
averaging. Internal Agent1 is the row axis and Agent2 the column axis; canonical
drug ordering always moves the corresponding dose. Responses are percent control
viability, with inhibition = 1 - response / 100. Values outside the inclusive
inhibition quality range [-2, 2] are preserved in exclusions rather than clipped.

No fitted statistics occur here. Cell/drug mappings, exclusions, source hashes,
and feature definitions are recorded before downstream split-specific fitting.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from concurrent.futures import ProcessPoolExecutor
from datetime import datetime, timezone
import hashlib
import importlib.metadata
import json
import logging
import math
import os
from pathlib import Path
import re
import shutil
import tempfile
import time
from typing import Any, Iterable
import urllib.error
import urllib.request
import warnings

import numpy as np
from numpy.typing import NDArray
from openpyxl import load_workbook
import pandas as pd
from rdkit import Chem, DataStructs, rdBase
from rdkit.Chem import rdFingerprintGenerator

LOGGER = logging.getLogger(__name__)
SCHEMA_VERSION = "source-v2.4"
INHIBITION_QUALITY_MIN = -2.0
INHIBITION_QUALITY_MAX = 2.0
INHIBITION_QUALITY_REASON = "response_outside_quality_range"
CHEMISTRY_COLUMN = "SMILES or PubChem ID"
SOURCE_DIRECTORIES = {
    "dream": "raw/az_dream/DREAM_Challenge_matrices_corrected",
    "oi": "raw/az_dream/DREAM_OI_matrices_corrected",
}
KNOWN_BIOLOGICS = {
    "ADAM17": "MEDI3622; monoclonal antibody",
    "CD19": "MT-103 / Blinatumomab; bispecific antibody",
    "PDGFR_Ab": "MEDI-575; monoclonal antibody",
    "TNFSF10": "TRAIL; protein ligand",
}
KNOWN_UNRELEASED = {"ATR_3", "PIP5K1", "SLC16A4"}
OBSERVATION_COLUMNS = [
    "observation_id",
    "matrix_id",
    "source",
    "source_file",
    "experiment",
    "cell_id",
    "drug_a",
    "drug_b",
    "dose_a_um",
    "dose_b_um",
    "inhibition",
    "response_percent",
    "grid_row",
    "grid_col",
    "matrix_cell_label",
    "filename_cell_label",
    "agent1",
    "agent2",
    "unit1",
    "unit2",
    "canonical_order_swapped",
]
EXCLUSION_COLUMNS = [
    "source",
    "source_file",
    "matrix_id",
    "reason",
    "detail",
    "grid_row",
    "grid_col",
    "response_percent",
    "inhibition",
]
DEFAULTS: dict[str, Any] = {
    "sources": ["dream"],
    "fingerprint_bits": 2048,
    "fingerprint_radius": 2,
    "include_zero_doses": True,
    "aliases": {},
    "unsupported_drug_policy": "exclude",
    "unresolved_cell_policy": "exclude",
    "missing_response_policy": "exclude",
    "max_matrices": None,
    "workers": 4,
    "resolve_pubchem": True,
    "pubchem_timeout_seconds": 30,
    "pubchem_retries": 3,
}


def response_quality_control() -> dict[str, Any]:
    """Return the fixed, manifest-safe response quality-control specification."""
    return {
        "field": "inhibition",
        "minimum": INHIBITION_QUALITY_MIN,
        "maximum": INHIBITION_QUALITY_MAX,
        "bounds": "inclusive",
        "action": "exclude_without_clipping",
        "reason": INHIBITION_QUALITY_REASON,
        "stage": "source preprocessing before dataset splitting",
    }


def _hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _json_hash(value: Any) -> str:
    return hashlib.sha256(
        json.dumps(value, sort_keys=True, separators=(",", ":"), allow_nan=False).encode()
    ).hexdigest()


def normalize_cell_name(value: Any) -> str:
    """Normalize supplied name fields only; this does not infer new aliases."""
    return re.sub(r"[^A-Z0-9]", "", str(value).upper())


def _clean(value: Any) -> str | None:
    if value is None or pd.isna(value):
        return None
    text = str(value).strip()
    return None if text.lower() in {"", "nan", "none", "null", "na", "n/a"} else text


def _cosmic(value: Any) -> str | None:
    text = _clean(value)
    if text is None:
        return None
    if not re.fullmatch(r"[0-9]+(?:\.0)?", text):
        # CMP includes entries with multiple identifiers; these are not an exact
        # unique COSMIC key and must never be split/guessed for this mapping.
        return None
    return text.removesuffix(".0")


def _required_columns(frame: pd.DataFrame, columns: Iterable[str], source: Path) -> None:
    missing = set(columns) - set(frame.columns)
    if missing:
        raise ValueError(f"{source}: missing columns {sorted(missing)}")


def map_cells(
    dataset_root: Path,
    aliases: dict[str, dict[str, str]] | None = None,
) -> tuple[pd.DataFrame, dict[str, str], dict[str, Any]]:
    """Map DREAM cells to CMP by exact unique COSMIC identifiers.

    Returns canonical cells, normalized name -> canonical cell, and provenance.
    Configured aliases must supply ``cell_id`` and a nonempty ``evidence`` string.
    Ambiguous source names and conflicting aliases raise instead of guessing.
    """
    dataset_root = Path(dataset_root)
    dream_path = dataset_root / "raw/az_dream_molecular/cell_info.csv"
    cmp_paths = sorted((dataset_root / "raw/cell_model_passports").glob("model_list_*.csv"))
    if len(cmp_paths) != 1:
        raise ValueError(f"Expected one immutable CMP model_list release, found {cmp_paths}")
    cmp_path = cmp_paths[0]
    dream = pd.read_csv(dream_path, dtype=str)
    cmp = pd.read_csv(cmp_path, dtype=str, low_memory=False)
    name_columns = ["Sanger.Name", "Alternative.Name", "CCLE.Name"]
    _required_columns(dream, [*name_columns, "COSMIC"], dream_path)
    _required_columns(cmp, ["model_id", "COSMIC_ID"], cmp_path)
    dream["cosmic_id"] = dream["COSMIC"].map(_cosmic)
    cmp["cosmic_id"] = cmp["COSMIC_ID"].map(_cosmic)
    if dream["cosmic_id"].isna().any() or dream["cosmic_id"].duplicated().any():
        raise ValueError("DREAM COSMIC IDs must be present and unique")
    relevant = cmp[cmp["cosmic_id"].isin(dream["cosmic_id"])]
    if relevant["cosmic_id"].duplicated().any():
        duplicates = relevant.loc[
            relevant["cosmic_id"].duplicated(False), ["cosmic_id", "model_id"]
        ]
        raise ValueError(f"Ambiguous CMP COSMIC mapping: {duplicates.to_dict('records')}")
    joined = dream.merge(
        relevant[["cosmic_id", "model_id"]],
        on="cosmic_id",
        how="left",
        validate="one_to_one",
    )
    if joined["model_id"].isna().any() or joined["model_id"].duplicated().any():
        raise ValueError("Every DREAM cell must map to a unique CMP model")
    joined["cell_id"] = joined["Sanger.Name"].map(_clean)
    if joined["cell_id"].isna().any() or joined["cell_id"].duplicated().any():
        raise ValueError("Canonical Sanger.Name values must be present and unique")
    names: dict[str, str] = {}
    name_evidence: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in joined.to_dict("records"):
        for column in name_columns:
            label = _clean(row[column])
            if label is None:
                continue
            normalized = normalize_cell_name(label)
            if normalized in names and names[normalized] != row["cell_id"]:
                raise ValueError(f"Ambiguous normalized cell name {label!r}")
            names[normalized] = row["cell_id"]
            name_evidence[normalized].append(
                {"cell_id": row["cell_id"], "column": column, "label": label}
            )
    for alias, spec in sorted((aliases or {}).items()):
        if not isinstance(spec, dict) or set(spec) != {"cell_id", "evidence"}:
            raise ValueError(f"Alias {alias!r} requires exactly cell_id and evidence")
        if spec["cell_id"] not in set(joined["cell_id"]) or not _clean(spec["evidence"]):
            raise ValueError(f"Alias {alias!r} has an unknown cell or missing evidence")
        normalized = normalize_cell_name(alias)
        if not normalized or (normalized in names and names[normalized] != spec["cell_id"]):
            raise ValueError(f"Alias {alias!r} conflicts with released cell metadata")
        names[normalized] = spec["cell_id"]
        name_evidence[normalized].append(
            {
                "cell_id": spec["cell_id"],
                "column": "configured_alias",
                "label": alias,
                "evidence": spec["evidence"],
            }
        )
    joined["mapping_method"] = "unique_exact_cosmic"
    joined["dream_source"] = str(dream_path.relative_to(dataset_root))
    joined["cmp_source"] = str(cmp_path.relative_to(dataset_root))
    cells = (
        joined[
            [
                "cell_id",
                "model_id",
                "cosmic_id",
                *name_columns,
                "mapping_method",
                "dream_source",
                "cmp_source",
            ]
        ]
        .sort_values("cell_id")
        .reset_index(drop=True)
    )
    provenance = {
        "source_files": [
            str(dream_path.relative_to(dataset_root)),
            str(cmp_path.relative_to(dataset_root)),
        ],
        "name_evidence": dict(name_evidence),
        "aliases": aliases or {},
    }
    LOGGER.info(
        "Mapped %d DREAM cells to CMP using unique exact COSMIC IDs (%d supplied aliases)",
        len(cells),
        len(aliases or {}),
    )
    return cells, names, provenance


def resolve_smiles(value: Any) -> tuple[str | None, str]:
    """Validate released SMILES; numeric PubChem IDs require an external mapping.

    Semicolon alternatives are accepted only if every parsed alternative has the
    same canonical isomeric SMILES. This verifies equality rather than choosing
    arbitrarily among potentially distinct compounds.
    """
    text = _clean(value)
    if text is None:
        return None, "missing_structure"
    if re.fullmatch(r"[0-9]+(?:\.0)?", text):
        return None, "pubchem_only_no_local_structure"
    alternatives = [part.strip() for part in text.split(";")]
    canonical = []
    with rdBase.BlockLogs():
        for part in alternatives:
            mol = Chem.MolFromSmiles(part)
            if mol is None or mol.GetNumAtoms() == 0:
                return None, "invalid_smiles"
            canonical.append(Chem.MolToSmiles(mol, canonical=True, isomericSmiles=True))
    if len(set(canonical)) != 1:
        return None, "ambiguous_semicolon_structures"
    return canonical[0], ("equivalent_semicolon_smiles" if len(canonical) > 1 else "valid_smiles")


def _pubchem_property(payload: bytes, cid: int) -> tuple[str, str]:
    """Validate that a property response describes exactly the requested CID."""
    try:
        data = json.loads(payload)
        properties = data["PropertyTable"]["Properties"]
    except (json.JSONDecodeError, KeyError, TypeError) as exc:
        raise ValueError(f"PubChem CID {cid}: invalid property-response schema") from exc
    if not isinstance(properties, list) or len(properties) != 1 or properties[0].get("CID") != cid:
        raise ValueError(f"PubChem CID {cid}: returned identifier is not the exact requested CID")
    # PUG REST renamed IsomericSMILES to SMILES; support both documented response
    # names, never the connectivity-only property when stereo is requested.
    for key in ["SMILES", "IsomericSMILES"]:
        value = properties[0].get(key)
        if isinstance(value, str):
            smiles, status = resolve_smiles(value)
            if smiles is None or status != "valid_smiles":
                raise ValueError(f"PubChem CID {cid}: returned SMILES cannot be validated by RDKit")
            return smiles, key
    raise ValueError(f"PubChem CID {cid}: response lacks a stereochemical SMILES property")


def resolve_pubchem_cid(
    cid: int,
    cache_root: Path,
    *,
    timeout_seconds: int = 30,
    retries: int = 3,
) -> tuple[str, dict[str, Any], list[Path]]:
    """Resolve one released numeric CID with a permanent verified offline cache.

    No name search, similar-compound search or surrogate selection is performed.
    A successful cache entry contains the exact response bytes and an immutable
    manifest recording the retrieval URL, SHA256, requested/returned CID and
    RDKit canonical SMILES. Existing entries are always verified, never refreshed.
    HTTP failures are fatal after bounded retries, not silently converted to a
    different scientific cohort. Only publicly released CIDs leave the machine.
    """
    if isinstance(cid, bool) or not isinstance(cid, int) or cid < 1:
        raise ValueError("PubChem CID must be a positive integer")
    cache_root = Path(cache_root)
    entry_dir = cache_root / str(cid)
    response_path, manifest_path = (
        entry_dir / "response.json",
        entry_dir / "manifest.json",
    )
    url = (
        f"https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/cid/{cid}/property/IsomericSMILES/JSON"
    )
    if entry_dir.exists():
        if not response_path.is_file() or not manifest_path.is_file():
            raise ValueError(f"Incomplete immutable PubChem cache entry: {entry_dir}")
        manifest = json.loads(manifest_path.read_text())
        payload = response_path.read_bytes()
        if (
            manifest.get("schema_version") != "pubchem-cid-v1"
            or manifest.get("requested_cid") != cid
            or manifest.get("url") != url
            or manifest.get("response_sha256") != hashlib.sha256(payload).hexdigest()
        ):
            raise ValueError(f"PubChem cache integrity mismatch for CID {cid}: {entry_dir}")
        smiles, property_name = _pubchem_property(payload, cid)
        if smiles != manifest.get("canonical_isomeric_smiles") or property_name != manifest.get(
            "response_property"
        ):
            raise ValueError(f"PubChem cached structure mismatch for CID {cid}")
        LOGGER.info("Verified offline PubChem CID %d cache %s", cid, entry_dir)
        return smiles, manifest, [response_path, manifest_path]
    LOGGER.info("Resolving released exact PubChem CID %d: %s", cid, url)
    request = urllib.request.Request(
        url,
        headers={
            "Accept": "application/json",
            "User-Agent": "reverse-dose/0.1 research-identifier-resolution",
        },
    )
    payload: bytes | None = None
    for attempt in range(retries):
        try:
            with urllib.request.urlopen(request, timeout=timeout_seconds) as response:
                payload = response.read()
            break
        except (urllib.error.URLError, TimeoutError, OSError) as exc:
            if isinstance(exc, urllib.error.HTTPError) and exc.code not in {
                429,
                500,
                502,
                503,
                504,
            }:
                raise RuntimeError(
                    f"PubChem exact CID {cid} failed with HTTP {exc.code}; use resolve_pubchem=false for an explicit reduced local-only cohort"
                ) from exc
            if attempt + 1 == retries:
                raise RuntimeError(
                    f"PubChem exact CID {cid} failed after {retries} attempts; no cache or cohort substitution performed"
                ) from exc
            LOGGER.warning(
                "PubChem CID %d request failed (%s); retry %d/%d",
                cid,
                exc,
                attempt + 2,
                retries,
            )
            time.sleep(min(2**attempt, 8))
    if payload is None:
        raise RuntimeError(f"PubChem exact CID {cid}: no response")
    smiles, property_name = _pubchem_property(payload, cid)
    manifest = {
        "schema_version": "pubchem-cid-v1",
        "requested_cid": cid,
        "returned_cid": cid,
        "url": url,
        "retrieved_at_utc": datetime.now(timezone.utc).isoformat(),
        "response_sha256": hashlib.sha256(payload).hexdigest(),
        "response_property": property_name,
        "canonical_isomeric_smiles": smiles,
        "rdkit_version": rdBase.rdkitVersion,
        "resolution_method": "exact numeric identifier supplied in released chemistry metadata",
    }
    cache_root.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=f".{cid}.staging-", dir=cache_root))
    try:
        (staging / "response.json").write_bytes(payload)
        (staging / "manifest.json").write_text(
            json.dumps(manifest, indent=2, sort_keys=True) + "\n"
        )
        os.rename(staging, entry_dir)
    except BaseException:
        shutil.rmtree(staging, ignore_errors=True)
        raise
    # Stay below the official five-requests-per-second ceiling even when an
    # endpoint answers unusually quickly. This occurs only for newly fetched CIDs.
    time.sleep(0.25)
    LOGGER.info(
        "Cached verified PubChem CID %d response SHA256=%s",
        cid,
        manifest["response_sha256"],
    )
    return smiles, manifest, [response_path, manifest_path]


def build_drug_features(
    dataset_root: Path,
    bits: int,
    radius: int,
    *,
    resolve_pubchem: bool = False,
    pubchem_timeout_seconds: int = 30,
    pubchem_retries: int = 3,
) -> tuple[pd.DataFrame, NDArray[np.str_], NDArray[np.float32], list[str]]:
    """Generate Morgan vectors from released SMILES or exact released PubChem CIDs.

    Locally parseable primary/supplemental SMILES always take precedence over
    remote identifier resolution. The returned source list includes any verified
    PubChem cache artifacts consumed alongside the immutable chemistry releases.
    """
    primary_path = dataset_root / "raw/az_dream/Drug_info_release.csv"
    supplement_path = (
        dataset_root
        / "raw/az_dream/Extra DrugChemistryData to release for Nat Comms FINAL RELEASE_200219.xlsx"
    )
    primary = pd.read_csv(primary_path, dtype=str)
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style.*")
        supplement = pd.read_excel(supplement_path, sheet_name="Drug_Chemistry_Data", dtype=str)
    frames = [(primary, primary_path), (supplement, supplement_path)]
    indexed: list[dict[str, dict[str, Any]]] = []
    for frame, path in frames:
        _required_columns(frame, ["ChallengeName", CHEMISTRY_COLUMN], path)
        frame["ChallengeName"] = frame["ChallengeName"].map(_clean)
        if frame["ChallengeName"].isna().any() or frame["ChallengeName"].duplicated().any():
            raise ValueError(f"{path}: drug ChallengeName values must be present and unique")
        indexed.append(frame.set_index("ChallengeName").to_dict("index"))
    records, identifiers, fingerprints = [], [], []
    consumed_sources = [str(path.relative_to(dataset_root)) for _, path in frames]
    generator = rdFingerprintGenerator.GetMorganGenerator(
        radius=radius, fpSize=bits, includeChirality=False
    )
    for name in sorted(set(indexed[0]) | set(indexed[1])):
        values = [_clean(frame.get(name, {}).get(CHEMISTRY_COLUMN)) for frame in indexed]
        candidates = [resolve_smiles(value) for value in values]
        selected = next(
            (index for index, (smiles, _) in enumerate(candidates) if smiles is not None),
            None,
        )
        canonical = candidates[selected][0] if selected is not None else None
        reason = (
            candidates[selected][1]
            if selected is not None
            else ";".join(sorted({reason for _, reason in candidates}))
        )
        pubchem_manifest = None
        pubchem_cid = None
        if (
            selected is None
            and resolve_pubchem
            and name not in KNOWN_BIOLOGICS
            and name not in KNOWN_UNRELEASED
        ):
            cid_source = next(
                (
                    index
                    for index, value in enumerate(values)
                    if value is not None and re.fullmatch(r"[0-9]+(?:\.0)?", value)
                ),
                None,
            )
            if cid_source is not None:
                pubchem_cid = int(values[cid_source].removesuffix(".0"))
                canonical, pubchem_manifest, cache_paths = resolve_pubchem_cid(
                    pubchem_cid,
                    dataset_root / "processed/drug_features/pubchem/v1",
                    timeout_seconds=pubchem_timeout_seconds,
                    retries=pubchem_retries,
                )
                consumed_sources.extend(str(path.relative_to(dataset_root)) for path in cache_paths)
                selected, reason = cid_source, "resolved_exact_released_pubchem_cid"
        if name in KNOWN_BIOLOGICS:
            canonical, selected, reason = (
                None,
                None,
                "known_biologic_without_small_molecule_structure",
            )
        elif name in KNOWN_UNRELEASED:
            canonical, selected, reason = (
                None,
                None,
                "unreleased_challenge_agent_structure",
            )
        record = {
            "drug_id": name,
            "supported": canonical is not None,
            "canonical_smiles": canonical,
            "structure_status": reason,
            "primary_value": values[0],
            "supplement_value": values[1],
            "primary_parse_status": candidates[0][1],
            "supplement_parse_status": candidates[1][1],
            "selected_source": (
                str(frames[selected][1].relative_to(dataset_root)) if selected is not None else None
            ),
            "released_drug_name": _clean(indexed[1].get(name, {}).get("Drug name")),
            "known_biologic_identity": KNOWN_BIOLOGICS.get(name),
            "fingerprint_bits": bits,
            "fingerprint_radius": radius,
            "include_chirality": False,
            "pubchem_cid": pubchem_cid,
            "pubchem_response_sha256": (
                pubchem_manifest["response_sha256"] if pubchem_manifest else None
            ),
            "pubchem_request_url": (pubchem_manifest["url"] if pubchem_manifest else None),
            "pubchem_cache_manifest": (
                f"processed/drug_features/pubchem/v1/{pubchem_cid}/manifest.json"
                if pubchem_manifest
                else None
            ),
        }
        records.append(record)
        if canonical is not None:
            mol = Chem.MolFromSmiles(canonical)
            vector = np.zeros(bits, dtype=np.float32)
            DataStructs.ConvertToNumpyArray(generator.GetFingerprint(mol), vector)
            identifiers.append(name)
            fingerprints.append(vector)
        else:
            LOGGER.warning("Drug excluded: %s (%s)", name, reason)
    if not fingerprints:
        raise ValueError("No released drug structures can be featurized")
    LOGGER.info(
        "Generated Morgan fingerprints for %d/%d released agents",
        len(identifiers),
        len(records),
    )
    return (
        pd.DataFrame(records),
        np.asarray(identifiers, dtype=np.str_),
        np.stack(fingerprints),
        sorted(set(consumed_sources)),
    )


def concentration_factor_um(unit: str) -> float:
    """Return exact recognized molar-unit -> micromolar conversion factors."""
    normalized = unit.strip().replace("\\", "").replace("μ", "u").replace("µ", "u")
    if normalized == "muM":
        normalized = "uM"
    factors = {
        "M": 1e6,
        "mM": 1e3,
        "uM": 1.0,
        "nM": 1e-3,
        "pM": 1e-6,
        "mol/L": 1e6,
        "mmol/L": 1e3,
        "umol/L": 1.0,
        "nmol/L": 1e-3,
        "pmol/L": 1e-6,
    }
    if normalized not in factors:
        raise ValueError(
            f"Unsupported concentration unit {unit!r}; no molecular-weight or mass-unit inference is allowed"
        )
    return factors[normalized]


def _number(value: Any, context: str) -> float:
    if isinstance(value, bool):
        raise ValueError(f"{context}: boolean is not a scientific numeric value")
    try:
        result = float(value)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"{context}: invalid numeric value {value!r}") from exc
    if not math.isfinite(result):
        raise ValueError(f"{context}: non-finite numeric value {value!r}")
    return result


def parse_corrected_matrix(path: Path) -> dict[str, Any]:
    """Read and validate every stored grid cell, including malformed XLSX dimensions.

    Returns internal labels, micromolar axes, finite measured responses and missing
    point coordinates. Grid indices refer to the original worksheet (zero-based).
    Missing points are represented explicitly; caller applies the configured
    missing-response policy. Unexpected schema/numeric/unit errors always raise.
    """
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style.*")
        workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        # Corrected releases incorrectly declare dimension A1. Reading according
        # to that declared range silently discards the entire response grid.
        declared_dimensions = worksheet.calculate_dimension()
        worksheet.reset_dimensions()
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        raise ValueError(f"{path}: empty worksheet")
    width = max(map(len, rows))
    rows = [tuple(row) + (None,) * (width - len(row)) for row in rows]
    first_column = [str(row[0]).strip().lower() for row in rows]
    first_row = [str(value).strip().lower() for value in rows[0]]
    if first_column.count("(=agent 1)") != 1 or first_row.count("(=agent 2)") != 1:
        raise ValueError(f"{path}: expected one Agent 1 row marker and Agent 2 column marker")
    row_end, col_end = first_column.index("(=agent 1)"), first_row.index("(=agent 2)")
    if row_end < 3 or col_end < 3:
        raise ValueError(f"{path}: both dose axes require at least two points")
    metadata: dict[str, str] = {}
    for row in rows[row_end + 1 :]:
        key = _clean(row[0])
        if key is None:
            continue
        if key in metadata:
            raise ValueError(f"{path}: duplicate metadata field {key}")
        value = _clean(row[1])
        if value is not None:
            metadata[key] = value
    for key in ["Agent1", "Agent2", "Unit1", "Unit2", "Title"]:
        if key not in metadata:
            raise ValueError(f"{path}: missing required metadata {key}")
    axes = []
    for index, raw_axis in enumerate(
        ([rows[r][0] for r in range(1, row_end)], list(rows[0][1:col_end])), start=1
    ):
        factor = concentration_factor_um(metadata[f"Unit{index}"])
        axis = [_number(value, f"{path}: Agent {index} dose axis") * factor for value in raw_axis]
        if any(value < 0 or not math.isfinite(value) for value in axis) or len(set(axis)) != len(
            axis
        ):
            raise ValueError(
                f"{path}: dose axis must contain unique finite nonnegative concentrations"
            )
        axes.append(axis)
    points, missing = [], []
    for row in range(1, row_end):
        for column in range(1, col_end):
            value = rows[row][column]
            if _clean(value) is None:
                missing.append((row, column))
            else:
                points.append(
                    (
                        row,
                        column,
                        _number(value, f"{path}: response row {row}, column {column}"),
                    )
                )
    return {
        "metadata": metadata,
        "dose1_um": axes[0],
        "dose2_um": axes[1],
        "points": points,
        "missing_points": missing,
        "declared_dimensions": declared_dimensions,
        "actual_shape": [len(rows), width],
    }


def _parse_job(job: tuple[str, str, str]) -> tuple[str, str, dict[str, Any]]:
    path, source, relative = job
    return source, relative, parse_corrected_matrix(Path(path))


def _effective_config(config: dict[str, Any]) -> dict[str, Any]:
    unknown = set(config) - set(DEFAULTS)
    if unknown:
        raise ValueError(f"Unknown source preprocessing configuration keys: {sorted(unknown)}")
    effective = {**DEFAULTS, **config}
    if (
        not effective["sources"]
        or set(effective["sources"]) - set(SOURCE_DIRECTORIES)
        or len(set(effective["sources"])) != len(effective["sources"])
    ):
        raise ValueError("sources must be a nonempty unique subset of ['dream', 'oi']")
    for key in ["fingerprint_bits", "workers"]:
        if (
            not isinstance(effective[key], int)
            or isinstance(effective[key], bool)
            or effective[key] < 1
        ):
            raise ValueError(f"{key} must be a positive integer")
    if (
        not isinstance(effective["fingerprint_radius"], int)
        or isinstance(effective["fingerprint_radius"], bool)
        or effective["fingerprint_radius"] < 0
    ):
        raise ValueError("fingerprint_radius must be a nonnegative integer")
    if not isinstance(effective["include_zero_doses"], bool) or not isinstance(
        effective["aliases"], dict
    ):
        raise ValueError("include_zero_doses must be boolean and aliases must be a mapping")
    if not isinstance(effective["resolve_pubchem"], bool):
        raise ValueError("resolve_pubchem must be boolean")
    for key in ["pubchem_timeout_seconds", "pubchem_retries"]:
        if (
            not isinstance(effective[key], int)
            or isinstance(effective[key], bool)
            or effective[key] < 1
        ):
            raise ValueError(f"{key} must be a positive integer")
    maximum = effective["max_matrices"]
    if maximum is not None and (
        not isinstance(maximum, int) or isinstance(maximum, bool) or maximum < 1
    ):
        raise ValueError("max_matrices must be null or a positive integer")
    for key in [
        "unsupported_drug_policy",
        "unresolved_cell_policy",
        "missing_response_policy",
    ]:
        if effective[key] not in {"exclude", "error"}:
            raise ValueError(f"{key} must be 'exclude' or 'error'")
    return effective


def _safe_output(dataset_root: Path, output_dir: Path) -> None:
    raw = (dataset_root / "raw").resolve()
    target = output_dir.resolve()
    if target == raw or target.is_relative_to(raw) or raw.is_relative_to(target):
        raise ValueError("Source preprocessing output must not overlap immutable dataset/raw")


def preprocess_sources(
    dataset_root: Path, config: dict[str, Any], output_dir: Path
) -> dict[str, Any]:
    """Build source tables and fingerprints atomically, returning their manifest.

    Existing artifacts are reused only after matching code/dependencies/config,
    every consumed source hash, and all generated output hashes. Otherwise fail
    without overwriting. All source paths in artifacts are dataset-root-relative.
    ``max_matrices`` is a deterministic systematic subset for engineering smoke
    checks; complete experiments must use null.
    """
    dataset_root, output_dir = Path(dataset_root).resolve(), Path(output_dir).resolve()
    _safe_output(dataset_root, output_dir)
    effective = _effective_config(config)
    LOGGER.info("Inspecting corrected dose matrices: sources=%s", effective["sources"])
    paths: list[tuple[str, Path]] = []
    source_counts = {}
    for source in effective["sources"]:
        selected_paths = sorted((dataset_root / SOURCE_DIRECTORIES[source]).rglob("*.xlsx"))
        if not selected_paths:
            raise ValueError(f"No corrected {source} XLSX matrices found")
        source_counts[source] = len(selected_paths)
        paths.extend((source, path) for path in selected_paths)
    paths.sort(key=lambda item: str(item[1].relative_to(dataset_root)))
    maximum = effective["max_matrices"]
    if maximum is not None and len(paths) > maximum:
        indices = (
            [0]
            if maximum == 1
            else [round(i * (len(paths) - 1) / (maximum - 1)) for i in range(maximum)]
        )
        paths = [paths[index] for index in indices]
        LOGGER.warning(
            "Engineering subset: selected %d of %d matrices systematically",
            len(paths),
            sum(source_counts.values()),
        )
    cells, names, cell_provenance = map_cells(dataset_root, effective["aliases"])
    drugs, drug_ids, drug_features, chemistry_sources = build_drug_features(
        dataset_root,
        effective["fingerprint_bits"],
        effective["fingerprint_radius"],
        resolve_pubchem=effective["resolve_pubchem"],
        pubchem_timeout_seconds=effective["pubchem_timeout_seconds"],
        pubchem_retries=effective["pubchem_retries"],
    )
    source_paths = sorted(
        set(
            [
                *(str(path.relative_to(dataset_root)) for _, path in paths),
                *cell_provenance["source_files"],
                *chemistry_sources,
                "raw/az_dream/Combenefit_dose_response_fitting_code.txt",
            ]
        )
    )
    LOGGER.info("Hashing %d consumed raw sources", len(source_paths))
    consumed_files = [
        {
            "path": relative,
            "sha256": _hash_file(dataset_root / relative),
            "bytes": (dataset_root / relative).stat().st_size,
        }
        for relative in source_paths
    ]
    raw_files = [entry for entry in consumed_files if entry["path"].startswith("raw/")]
    resolved_chemistry_files = [
        entry for entry in consumed_files if not entry["path"].startswith("raw/")
    ]
    versions = {
        name: importlib.metadata.version(name)
        for name in ["numpy", "pandas", "openpyxl", "rdkit", "pyarrow"]
    }
    identity = {
        "schema_version": SCHEMA_VERSION,
        "config": effective,
        "response_quality_control": response_quality_control(),
        "raw_files": raw_files,
        "resolved_chemistry_files": resolved_chemistry_files,
        "source_inventory": source_counts,
        "source_code_sha256": _hash_file(Path(__file__)),
        "dependencies": versions,
    }
    identity_hash = _json_hash(identity)
    manifest_path = output_dir / "source_manifest.json"
    if output_dir.exists():
        if not manifest_path.is_file():
            raise FileExistsError(f"Output exists without a complete source manifest: {output_dir}")
        manifest_checksum_path = output_dir / "source_manifest.sha256"
        if (
            not manifest_checksum_path.is_file()
            or manifest_checksum_path.read_text().strip() != _hash_file(manifest_path)
        ):
            raise ValueError(f"Source manifest integrity mismatch: {manifest_path}")
        cached = json.loads(manifest_path.read_text())
        if cached.get("identity_hash") != identity_hash or any(
            cached.get(key) != value for key, value in identity.items()
        ):
            raise ValueError(
                f"Existing source artifacts differ from raw/config/code/dependencies; use a new output directory: {output_dir}"
            )
        required_outputs = {
            "cells.parquet",
            "drugs.parquet",
            "observations.parquet",
            "exclusions.parquet",
            "drugs.npz",
        }
        if (
            not isinstance(cached.get("output_hashes"), dict)
            or set(cached["output_hashes"]) != required_outputs
        ):
            raise ValueError(
                f"Source output hash inventory is incomplete or unexpected: {output_dir}"
            )
        for relative, digest in cached["output_hashes"].items():
            if _hash_file(output_dir / relative) != digest:
                raise ValueError(f"Cached artifact integrity mismatch: {relative}")
        LOGGER.info("Reusing verified source cache %s (%s)", output_dir, identity_hash[:12])
        return cached
    observations: list[dict[str, Any]] = []
    exclusions: list[dict[str, Any]] = []
    supported = set(drug_ids.tolist())
    known_names = {str(name).casefold(): str(name) for name in drugs["drug_id"]}
    if len(known_names) != len(drugs):
        raise ValueError("Drug IDs are ambiguous under case-insensitive matching")
    # The release explicitly lists this observed alternate spelling. Verify the
    # exact token and canonical ID before using it. Do not generically reinterpret
    # comma-separated drug names: some describe multi-agent cocktail constituents.
    drug_alias_evidence = {}
    gemcitabine = drugs.loc[drugs["drug_id"] == "Gemcitabine"]
    if not gemcitabine.empty:
        drug_name = str(gemcitabine.iloc[0]["released_drug_name"])
        if "Gemcitibine" in [token.strip() for token in drug_name.split(",")]:
            known_names["gemcitibine"] = "Gemcitabine"
            drug_alias_evidence["Gemcitibine"] = {
                "drug_id": "Gemcitabine",
                "source": "raw/az_dream/Extra DrugChemistryData to release for Nat Comms FINAL RELEASE_200219.xlsx",
                "sheet": "Drug_Chemistry_Data",
                "column": "Drug name",
                "released_value": drug_name,
            }
    malformed_dimensions = 0
    accepted_matrices = 0
    jobs = [(str(path), source, str(path.relative_to(dataset_root))) for source, path in paths]
    executor = (
        ProcessPoolExecutor(max_workers=effective["workers"]) if effective["workers"] > 1 else None
    )
    try:
        parsed = (
            executor.map(_parse_job, jobs, chunksize=32)
            if executor is not None
            else map(_parse_job, jobs)
        )
        for index, (source, relative, matrix) in enumerate(parsed, start=1):
            matrix_id = hashlib.sha256(relative.encode()).hexdigest()[:24]
            metadata = matrix["metadata"]
            filename = Path(relative).stem.split("&")
            if len(filename) != 4:
                raise ValueError(f"{relative}: filename must have DrugA&DrugB&Cell&Experiment")
            if matrix["declared_dimensions"] in {"A1", "A1:A1"} and matrix["actual_shape"] != [
                1,
                1,
            ]:
                malformed_dimensions += 1
            reasons: list[tuple[str, str]] = []
            raw_agent1, raw_agent2 = metadata["Agent1"], metadata["Agent2"]
            filename_drugs = [known_names.get(label.casefold(), label) for label in filename[:2]]
            agent1 = known_names.get(raw_agent1.casefold(), raw_agent1)
            agent2 = known_names.get(raw_agent2.casefold(), raw_agent2)
            if sorted(filename_drugs) != sorted([agent1, agent2]):
                raise ValueError(f"{relative}: filename and internal drug labels disagree")
            unsupported = sorted({agent1, agent2} - supported)
            if unsupported:
                reasons.append(("unsupported_drug_structure", json.dumps(unsupported)))
                if effective["unsupported_drug_policy"] == "error":
                    raise ValueError(f"{relative}: unsupported drug structures {unsupported}")
            title_cell, filename_cell = (
                names.get(normalize_cell_name(metadata["Title"])),
                names.get(normalize_cell_name(filename[2])),
            )
            if title_cell is None or filename_cell is None:
                reasons.append(
                    (
                        "unresolved_cell_label",
                        json.dumps({"title": metadata["Title"], "filename": filename[2]}),
                    )
                )
            if normalize_cell_name(metadata["Title"]) != normalize_cell_name(filename[2]) and (
                title_cell is None or filename_cell is None or title_cell != filename_cell
            ):
                reasons.append(
                    (
                        "cell_label_mismatch",
                        json.dumps({"title": metadata["Title"], "filename": filename[2]}),
                    )
                )
            elif (
                title_cell is not None and filename_cell is not None and title_cell != filename_cell
            ):
                reasons.append(
                    (
                        "cell_label_mismatch",
                        json.dumps({"title": metadata["Title"], "filename": filename[2]}),
                    )
                )
            if effective["unresolved_cell_policy"] == "error" and any(
                reason.startswith(("unresolved_cell", "cell_label")) for reason, _ in reasons
            ):
                raise ValueError(f"{relative}: unresolved or conflicting cell labels {reasons}")
            if matrix["missing_points"] and effective["missing_response_policy"] == "error":
                raise ValueError(f"{relative}: missing response points {matrix['missing_points']}")
            if not matrix["points"]:
                reasons.append(
                    (
                        "no_measured_responses",
                        "Every response in the released dose grid is missing; no labels constructed",
                    )
                )
            for row, column in matrix["missing_points"]:
                exclusions.append(
                    dict(
                        source=source,
                        source_file=relative,
                        matrix_id=matrix_id,
                        reason="missing_response",
                        detail="No released finite response; no imputation",
                        grid_row=row,
                        grid_col=column,
                    )
                )
            for reason, detail in reasons:
                exclusions.append(
                    dict(
                        source=source,
                        source_file=relative,
                        matrix_id=matrix_id,
                        reason=reason,
                        detail=detail,
                        grid_row=None,
                        grid_col=None,
                    )
                )
            if not reasons:
                before = len(observations)
                swapped = agent1 > agent2
                drug_a, drug_b = sorted([agent1, agent2])
                for row, column, response in matrix["points"]:
                    dose1, dose2 = (
                        matrix["dose1_um"][row - 1],
                        matrix["dose2_um"][column - 1],
                    )
                    inhibition = 1.0 - response / 100.0
                    if not INHIBITION_QUALITY_MIN <= inhibition <= INHIBITION_QUALITY_MAX:
                        exclusions.append(
                            dict(
                                source=source,
                                source_file=relative,
                                matrix_id=matrix_id,
                                reason=INHIBITION_QUALITY_REASON,
                                detail=(
                                    "Computed inhibition is outside inclusive quality range "
                                    f"[{INHIBITION_QUALITY_MIN}, {INHIBITION_QUALITY_MAX}]; "
                                    "raw response retained without clipping"
                                ),
                                grid_row=row,
                                grid_col=column,
                                response_percent=response,
                                inhibition=inhibition,
                            )
                        )
                        continue
                    if not effective["include_zero_doses"] and (dose1 == 0 or dose2 == 0):
                        exclusions.append(
                            dict(
                                source=source,
                                source_file=relative,
                                matrix_id=matrix_id,
                                reason="zero_dose_excluded_by_config",
                                detail="At least one dose coordinate equals zero",
                                grid_row=row,
                                grid_col=column,
                            )
                        )
                        continue
                    dose_a, dose_b = (dose2, dose1) if swapped else (dose1, dose2)
                    observations.append(
                        dict(
                            observation_id=f"{matrix_id}:{row}:{column}",
                            matrix_id=matrix_id,
                            source=source,
                            source_file=relative,
                            experiment=filename[3],
                            cell_id=title_cell,
                            drug_a=drug_a,
                            drug_b=drug_b,
                            dose_a_um=dose_a,
                            dose_b_um=dose_b,
                            inhibition=inhibition,
                            response_percent=response,
                            grid_row=row,
                            grid_col=column,
                            matrix_cell_label=metadata["Title"],
                            filename_cell_label=filename[2],
                            agent1=raw_agent1,
                            agent2=raw_agent2,
                            unit1=metadata["Unit1"],
                            unit2=metadata["Unit2"],
                            canonical_order_swapped=swapped,
                        )
                    )
                accepted_matrices += int(len(observations) > before)
            if index % 500 == 0 or index == len(paths):
                LOGGER.info(
                    "Parsed %d/%d matrices; retained %d matrices, %d observations; %d exclusions",
                    index,
                    len(paths),
                    accepted_matrices,
                    len(observations),
                    len(exclusions),
                )
    finally:
        if executor is not None:
            executor.shutdown(wait=True, cancel_futures=True)
    if not observations:
        raise ValueError("No observations remain after explicit source inclusion policies")
    observation_frame = pd.DataFrame(observations, columns=OBSERVATION_COLUMNS)
    if observation_frame["observation_id"].duplicated().any():
        raise ValueError("Generated observation IDs must be unique")
    exclusion_frame = pd.DataFrame(exclusions, columns=EXCLUSION_COLUMNS)
    manifest = {
        **identity,
        "identity_hash": identity_hash,
        "cell_mapping": cell_provenance,
        "observation_schema": OBSERVATION_COLUMNS,
        "exclusion_schema": EXCLUSION_COLUMNS,
        "drug_alias_evidence": drug_alias_evidence,
        "response_definition": (
            "inhibition = 1 - percent_control_viability / 100; no clipping; "
            "retain only inclusive quality range [-2.0, 2.0]"
        ),
        "response_source_evidence": "raw/az_dream/Combenefit_dose_response_fitting_code.txt fixes EC0=100.0",
        "dose_unit": "uM",
        "dose_order": "case-sensitive lexicographic canonical drug IDs; doses moved together",
        "fingerprint_definition": {
            "algorithm": "RDKit Morgan bit vector",
            "radius": effective["fingerprint_radius"],
            "bits": effective["fingerprint_bits"],
            "includeChirality": False,
            "exact_released_pubchem_resolution": effective["resolve_pubchem"],
            "semicolon_policy": "accept only canonical isomeric equality of all alternatives",
            "source_precedence": "parseable primary, then parseable supplement, then exact released numeric CID (primary takes precedence)",
        },
        "counts": {
            "source_matrices": source_counts,
            "parsed_matrices": len(paths),
            "retained_matrices": accepted_matrices,
            "retained_observations": len(observation_frame),
            "mapped_cells": len(cells),
            "retained_cells": int(observation_frame["cell_id"].nunique()),
            "supported_drugs": len(drug_ids),
            "retained_drugs": len(
                set(observation_frame["drug_a"]) | set(observation_frame["drug_b"])
            ),
            "incorrect_declared_dimensions_read_completely": malformed_dimensions,
            "excluded_by_reason": dict(sorted(Counter(exclusion_frame["reason"]).items())),
            "quality_control_excluded_observations": int(
                (exclusion_frame["reason"] == INHIBITION_QUALITY_REASON).sum()
            ),
            "retained_observations_outside_unit_interval": int(
                (
                    (observation_frame["inhibition"] < 0) | (observation_frame["inhibition"] > 1)
                ).sum()
            ),
        },
    }
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=f".{output_dir.name}.staging-", dir=output_dir.parent))
    try:
        cells.to_parquet(staging / "cells.parquet", index=False)
        drugs.to_parquet(staging / "drugs.parquet", index=False)
        observation_frame.to_parquet(staging / "observations.parquet", index=False)
        exclusion_frame.to_parquet(staging / "exclusions.parquet", index=False)
        np.savez_compressed(
            staging / "drugs.npz",
            ids=drug_ids,
            features=drug_features.astype(np.float32, copy=False),
        )
        manifest["output_hashes"] = {
            path.name: _hash_file(path) for path in sorted(staging.iterdir())
        }
        (staging / "source_manifest.json").write_text(
            json.dumps(manifest, indent=2, sort_keys=True, allow_nan=False) + "\n"
        )
        (staging / "source_manifest.sha256").write_text(
            _hash_file(staging / "source_manifest.json") + "\n"
        )
        os.rename(staging, output_dir)
    except BaseException:
        shutil.rmtree(staging, ignore_errors=True)
        raise
    LOGGER.info(
        "Source artifacts committed to %s; identity=%s; exclusions=%s",
        output_dir,
        identity_hash,
        manifest["counts"]["excluded_by_reason"],
    )
    return manifest
