"""Read official ALMANAC and O'Neil measurements without fitting any features."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
import pandas as pd

from .preprocess import concentration_factor_um, normalize_cell_name


def one_file(directory: Path, pattern: str) -> Path:
    files = sorted(directory.glob(pattern))
    if len(files) != 1:
        raise ValueError(f"Expected exactly one {pattern} in {directory}, found {len(files)}")
    return files[0]


def _require(frame: pd.DataFrame, columns: set[str]) -> None:
    if columns - set(frame):
        raise ValueError(f"Missing required released columns: {sorted(columns - set(frame))}")
    if frame.empty:
        raise ValueError("Released response table is empty")


def _matrix_ids(frame: pd.DataFrame, columns: list[str], source_file: str) -> pd.Series:
    """Hash full experiment keys, never row position, plate alone, or drug pair alone."""
    keys = frame[columns].astype(str)
    index = pd.MultiIndex.from_frame(keys)
    codes, uniques = pd.factorize(index, sort=True)
    names = np.asarray(
        [
            hashlib.sha256(
                json.dumps([source_file, *key], separators=(",", ":")).encode()
            ).hexdigest()[:24]
            for key in uniques
        ]
    )
    return pd.Series(names[codes], index=frame.index)


def _coordinates(frame: pd.DataFrame) -> pd.DataFrame:
    """Use actual concentrations, including source index corrections, to define axes."""
    doses = frame[["dose1_um", "dose2_um"]].to_numpy(dtype=float)
    if not np.isfinite(doses).all() or (doses < 0).any():
        raise ValueError("Released dose coordinates must be finite and nonnegative")
    group = frame.groupby("matrix_id", sort=False, observed=True)
    frame["grid_row"] = group.dose1_um.rank(method="dense").astype(np.int32)
    frame["grid_col"] = group.dose2_um.rank(method="dense").astype(np.int32)
    if frame.duplicated(["matrix_id", "grid_row", "grid_col"]).any():
        raise ValueError("Duplicate coordinates within a released experiment/replicate")
    shape = frame.groupby("matrix_id", observed=True).agg(
        rows=("grid_row", "max"), columns=("grid_col", "max"), points=("grid_row", "size")
    )
    if shape["points"].ne(shape["rows"] * shape["columns"]).any():
        raise ValueError("Released experiment has an incomplete dose grid")
    frame["observation_id"] = (
        frame.matrix_id + ":" + frame.grid_row.astype(str) + ":" + frame.grid_col.astype(str)
    )
    return frame


def read_oneil(
    path: Path, source_file: str, config: dict[str, Any]
) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Preserve batch and all four technical replicates, including missing points."""
    with path.open("rb") as handle:
        workbook = load_workbook(handle, read_only=True, data_only=True)
        try:
            if workbook.sheetnames != ["combination response"]:
                raise ValueError("Expected the official O'Neil combination-response worksheet")
            rows = iter(workbook.active.values)
            frame = pd.DataFrame(rows, columns=next(rows))
        finally:
            workbook.close()
    required = {
        "BatchID",
        "cell_line",
        "drugA_name",
        "drugB_name",
        "drugA Conc (µM)",
        "drugB Conc (µM)",
        *[f"viability{i}" for i in range(1, 5)],
    }
    _require(frame, required)
    for col in ("BatchID", "cell_line", "drugA_name", "drugB_name"):
        if frame[col].isna().any():
            raise ValueError(f"Missing O'Neil {col}")
    if not frame.BatchID.isin([1, 2, 3]).all():
        raise ValueError("Unknown O'Neil batch identifier")
    total_rows = len(frame)
    frame["source_row"] = np.arange(2, len(frame) + 2)
    frame = frame.loc[frame.BatchID.isin(config["oneil_batches"])].copy()
    frame["drug1"], frame["drug2"] = frame.drugA_name.astype(str), frame.drugB_name.astype(str)
    frame["cell_label"] = frame.cell_line.astype(str)
    frame["dose1_um"] = pd.to_numeric(frame["drugA Conc (µM)"], errors="raise")
    frame["dose2_um"] = pd.to_numeric(frame["drugB Conc (µM)"], errors="raise")
    frame["experiment"] = "batch:" + frame.BatchID.astype(str)
    base_ids = _matrix_ids(frame, ["BatchID", "cell_line", "drug1", "drug2"], source_file)
    columns = ["cell_label", "drug1", "drug2", "dose1_um", "dose2_um", "experiment", "source_row"]
    parts = []
    for replicate in range(1, 5):
        part = frame[columns].copy()
        part["matrix_id"] = base_ids + f"-r{replicate}"
        values = frame[f"viability{replicate}"].replace("NULL", np.nan)
        # Only the publisher's missing marker is allowed; malformed text is not
        # silently turned into missing response data.
        part["response_percent"] = pd.to_numeric(values, errors="raise") * 100
        part["response_origin"] = f"viability{replicate}"
        part["replicate"] = replicate
        parts.append(part)
    result = _coordinates(pd.concat(parts, ignore_index=True))
    return result, {
        "released_dose_pair_rows": total_rows,
        "selected_dose_pair_rows": len(frame),
        "physical_experiments": int(base_ids.nunique()),
        "technical_replicates": 4,
        "response_definition": "inhibition = 1 - released fractional viability; technical replicates retained separately",
        "response_source_evidence": "publisher/22503130.json; official response legend: viability1..4 are technical replicates",
        "monotherapy_policy": "Separate single-agent supplement is archived as reference; no unmeasured pair-specific zero-dose axes are constructed",
    }


def read_almanac(
    path: Path, source_file: str, config: dict[str, Any]
) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Read measured pairs; reconstruct missing no-TZ only under an explicit policy."""
    text = [
        "SCREENER",
        "STUDY",
        "TESTDATE",
        "PLATE",
        "PREFIX1",
        "PREFIX2",
        "CONCUNIT1",
        "CONCUNIT2",
        "VALID",
        "PANEL",
        "CELLNAME",
    ]
    frame = pd.read_csv(path, dtype={c: "category" for c in text}, low_memory=False)
    _require(
        frame,
        {
            "COMBODRUGSEQ",
            "SCREENER",
            "STUDY",
            "TESTDATE",
            "PANELNBR",
            "CELLNBR",
            "NSC1",
            "SAMPLE1",
            "NSC2",
            "SAMPLE2",
            "CONC1",
            "CONC2",
            "CONCUNIT1",
            "CONCUNIT2",
            "VALID",
            "CELLNAME",
            "PERCENTGROWTH",
            "PERCENTGROWTHNOTZ",
            "TESTVALUE",
            "CONTROLVALUE",
        },
    )
    if frame.COMBODRUGSEQ.isna().any() or frame.COMBODRUGSEQ.duplicated().any():
        raise ValueError("ALMANAC record identifiers must be present and unique")
    single_rows = int(frame.NSC2.isna().sum())
    frame = frame.loc[frame.NSC2.notna()].copy()
    # CELLNBR is only unique within PANELNBR. Preserve the pair and validate all
    # label variants (including the released SF-539 control character).
    labels = frame.CELLNAME.astype(str).map(normalize_cell_name)
    if (
        labels.eq("").any()
        or labels.groupby([frame.PANELNBR, frame.CELLNBR], observed=True).nunique().gt(1).any()
    ):
        raise ValueError("Conflicting ALMANAC cell labels for a panel/cell identifier")
    for col in (
        "SCREENER",
        "STUDY",
        "TESTDATE",
        "PANELNBR",
        "CELLNBR",
        "NSC1",
        "SAMPLE1",
        "SAMPLE2",
        "CELLNAME",
    ):
        if frame[col].isna().any():
            raise ValueError(f"Missing ALMANAC {col}")
    for col in ("NSC1", "NSC2"):
        if (frame[col] <= 0).any() or (frame[col] != frame[col].astype(np.int64)).any():
            raise ValueError("NSC identifiers must be positive integers")
    frame["drug1"] = "NSC" + frame.NSC1.astype(np.int64).astype(str)
    frame["drug2"] = "NSC" + frame.NSC2.astype(np.int64).astype(str)
    for number in (1, 2):
        units = frame[f"CONCUNIT{number}"]
        if units.isna().any():
            raise ValueError("Missing ALMANAC combination concentration unit")
        factors = units.astype(str).map(
            {unit: concentration_factor_um(unit) for unit in units.unique()}
        )
        frame[f"unit{number}"] = units.astype(str)
        frame[f"dose{number}_um"] = pd.to_numeric(frame[f"CONC{number}"], errors="raise") * factors
    keys = [
        "SCREENER",
        "STUDY",
        "TESTDATE",
        "PANELNBR",
        "CELLNBR",
        "NSC1",
        "SAMPLE1",
        "NSC2",
        "SAMPLE2",
    ]
    # 5x3 grids span several plates; PLATE remains row provenance, not a matrix key.
    frame["matrix_id"] = _matrix_ids(frame, keys, source_file)
    frame["cell_label"] = frame.CELLNAME.astype(str)
    frame["experiment"] = (
        frame.SCREENER.astype(str)
        + ":"
        + frame.STUDY.astype(str)
        + ":"
        + frame.TESTDATE.astype(str)
    )
    frame["source_row"] = frame.COMBODRUGSEQ.astype(np.int64)
    missing = frame.PERCENTGROWTHNOTZ.isna()
    if config["nci_response"] == "percent_growth":
        response = frame.PERCENTGROWTH.copy()
        origin = pd.Series("reported_percent_growth", index=frame.index)
        definition = "inhibition = 1 - released time-zero-corrected PERCENTGROWTH / 100; growth endpoint, not percent-control viability"
    else:
        response = frame.PERCENTGROWTHNOTZ.copy()
        origin = pd.Series("reported_no_tz", index=frame.index)
        if missing.any() and config["nci_missing_no_tz"] == "error":
            raise ValueError("Missing ALMANAC PERCENTGROWTHNOTZ under error policy")
        if config["nci_missing_no_tz"] == "reconstruct":
            valid = (
                frame.CONTROLVALUE.gt(0)
                & np.isfinite(frame.CONTROLVALUE)
                & np.isfinite(frame.TESTVALUE)
            )
            response.loc[missing & valid] = (
                100
                * frame.loc[missing & valid, "TESTVALUE"]
                / frame.loc[missing & valid, "CONTROLVALUE"]
            )
            origin.loc[missing & valid] = "reconstructed_test_over_control"
        definition = "inhibition = 1 - percent-control response / 100; PERCENTGROWTHNOTZ or explicitly reconstructed 100*TESTVALUE/CONTROLVALUE; never substitute time-zero-corrected growth"
    frame["response_percent"], frame["response_origin"] = response, origin
    frame["replicate"] = 1
    frame["valid_source"] = frame.VALID.eq("Y")
    result = _coordinates(frame)
    keep = [
        "matrix_id",
        "observation_id",
        "cell_label",
        "drug1",
        "drug2",
        "dose1_um",
        "dose2_um",
        "grid_row",
        "grid_col",
        "response_percent",
        "response_origin",
        "experiment",
        "source_row",
        "replicate",
        "valid_source",
        "unit1",
        "unit2",
    ]
    return result[keep].reset_index(drop=True), {
        "single_agent_rows_not_paired": single_rows,
        "released_combination_rows": len(result),
        "physical_experiments": int(result.matrix_id.nunique()),
        "missing_reported_no_tz": int(missing.sum()),
        "reconstructed_no_tz": int(origin.eq("reconstructed_test_over_control").sum()),
        "response_definition": definition,
        "response_source_evidence": "ALMANAC_DataFields.txt: PercentGrowth uses time zero; PercentGrowthNoTZ does not",
        "monotherapy_policy": "Standalone single-agent records have no second drug; not duplicated into pair-specific zero-dose axes",
    }
